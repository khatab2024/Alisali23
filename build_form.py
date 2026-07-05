# -*- coding: utf-8 -*-
"""
استمارة الفحص الفني الشامل - قسم الدراسات والبحوث (ER-0001)
نسخة Excel ذكية مطابقة للأصل مع أتمتة كاملة:
- ترويسة رسمية بالشعار + رقم الاستمارة + تاريخ الإصدار
- بيانات الاستمارة الأساسية
- تفاصيل الفحص (م1..م30) مع قوائم منسدلة وبحث تلقائي وتظليل ذكي
- ملخص النتائج (COUNTIF + نسبة القبول)
- تواقيع: الضبط والجودة / مدير المشروع أو المنشأة
- ورقة "البيانات" لتعريف القطع وعدد مراحلها
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.drawing.image import Image as XLImage

# ---------------- ألوان وخطوط ----------------
NAVY, BLUE, LIGHT = "1F3864", "2E5496", "D9E1F2"
GREY, GREEN, RED, WHITE, GOLD = "BFBFBF", "C6EFCE", "FFC7CE", "FFFFFF", "BF9000"

thin = Side(style="thin", color="808080")
med  = Side(style="medium", color="404040")
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
border_box = Border(left=med, right=med, top=med, bottom=med)

def font(size=11, bold=False, color="000000"):
    return Font(name="Arial", size=size, bold=bold, color=color)

center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right  = Alignment(horizontal="right",  vertical="center", wrap_text=True)

def fill(c):
    return PatternFill("solid", fgColor=c)

# ---------------- بيانات القطع (عيّنة قابلة للتعديل) ----------------
# رقم القطعة | اسم القطعة | عدد المراحل
PARTS = [
    ("P-1001", "عمود إدارة رئيسي", 10),
    ("P-1002", "غطاء علبة التروس", 20),
    ("P-1003", "جسم المضخة",        30),
    ("P-1004", "فلنجة توصيل",       10),
    ("P-1005", "ذراع رافعة",        20),
    ("P-1006", "قاعدة محرك",        30),
]
MAX_STAGES = 30
DATA_MAX_ROW = 200
LOGO = "/projects/sandbox/logo.png"

wb = Workbook()

# ==================================================================
# ورقة البيانات (مرجع الأتمتة)
# ==================================================================
wsd = wb.active
wsd.title = "البيانات"
wsd.sheet_view.rightToLeft = True
for c, h in enumerate(["رقم القطعة", "اسم القطعة", "عدد المراحل"], start=1):
    cell = wsd.cell(row=1, column=c, value=h)
    cell.font = font(12, bold=True, color=WHITE)
    cell.fill = fill(NAVY); cell.alignment = center; cell.border = border_all
for i, (num, name, st) in enumerate(PARTS, start=2):
    wsd.cell(row=i, column=1, value=num).alignment = center
    wsd.cell(row=i, column=2, value=name).alignment = right
    wsd.cell(row=i, column=3, value=st).alignment = center
    for c in range(1, 4):
        wsd.cell(row=i, column=c).border = border_all
        wsd.cell(row=i, column=c).font = font(11)
wsd.column_dimensions["A"].width = 16
wsd.column_dimensions["B"].width = 30
wsd.column_dimensions["C"].width = 14
n = wsd.cell(row=len(PARTS)+3, column=1,
             value="أضف/عدّل القطع هنا (رقم القطعة، اسمها، عدد مراحلها). تنعكس تلقائياً في الاستمارة.")
n.font = font(10, bold=True, color="C00000"); n.alignment = right
wsd.merge_cells(start_row=len(PARTS)+3, start_column=1, end_row=len(PARTS)+3, end_column=3)

wb.defined_names["PartsList"]  = DefinedName("PartsList",  attr_text=f"'البيانات'!$A$2:$A${DATA_MAX_ROW}")
wb.defined_names["PartsTable"] = DefinedName("PartsTable", attr_text=f"'البيانات'!$A$2:$C${DATA_MAX_ROW}")

# ==================================================================
# ورقة الاستمارة
# ==================================================================
ws = wb.create_sheet("استمارة الفحص")
ws.sheet_view.rightToLeft = True
ws.sheet_view.showGridLines = False

widths = {"A": 6, "B": 28, "C": 12, "D": 20, "E": 13, "F": 13, "G": 28}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# ---------------- الترويسة الرسمية ----------------
# شعار الجهة (يمين) + رقم الاستمارة وتاريخ الإصدار (يسار)
ws.merge_cells("A1:B3")
gh = ws["A1"]
gh.value = "قسم الدراسات والبحوث\n(شعار الجهة)"
gh.font = font(12, bold=True, color=NAVY)
gh.alignment = center

ws.merge_cells("F1:G1"); ws["F1"].value = "رقم الاستمارة:  ER-0001"
ws.merge_cells("F2:G2"); ws["F2"].value = "تاريخ الإصدار:  ..../..../...."
ws.merge_cells("F3:G3"); ws["F3"].value = "رقم الإصدار:  (1)"
for r in (1, 2, 3):
    ws[f"F{r}"].font = font(10, bold=True, color=NAVY)
    ws[f"F{r}"].alignment = right

# الشعار في الوسط
if os.path.exists(LOGO):
    img = XLImage(LOGO)
    img.width, img.height = 95, 80
    ws.add_image(img, "D1")
for r in (1, 2, 3):
    ws.row_dimensions[r].height = 22

# ---------------- العنوان ----------------
ws.merge_cells("A4:G4")
t = ws["A4"]
t.value = "استمارة الفحص الفني الشامل"
t.font = font(20, bold=True, color=WHITE)
t.fill = fill(NAVY); t.alignment = center
ws.row_dimensions[4].height = 40

# ================= القسم 1: بيانات الاستمارة الأساسية =================
ws.merge_cells("A5:G5")
s1 = ws["A5"]
s1.value = "◆  بيانات الاستمارة الأساسية"
s1.font = font(13, bold=True, color=WHITE); s1.fill = fill(BLUE); s1.alignment = right
ws.row_dimensions[5].height = 26

def label(ref, text):
    c = ws[ref]; c.value = text
    c.font = font(11, bold=True, color=NAVY); c.fill = fill(LIGHT)
    c.alignment = right; c.border = border_all

def valuecell(ref, formula=None):
    c = ws[ref]
    if formula:
        c.value = formula
    c.font = font(11, bold=True); c.alignment = center
    c.border = border_all; c.fill = fill(WHITE)

# صف 6: رقم القطعة | المجموعة
label("A6", "رقم القطعة"); valuecell("B6")                     # قائمة منسدلة
label("C6", "المجموعة");   ws.merge_cells("D6:G6"); valuecell("D6")
# صف 7: اسم القطعة | التاريخ
label("A7", "اسم القطعة")
ws.merge_cells("B7:C7"); valuecell("B7", "=IFERROR(VLOOKUP($B$6,PartsTable,2,FALSE),\"\")")
label("D7", "التاريخ");    ws.merge_cells("E7:G7"); valuecell("E7")
ws["E7"].value = "..../..../...."
# صف 8: اسم المشروع | عدد المراحل
label("A8", "اسم المشروع")
ws.merge_cells("B8:D8"); valuecell("B8")
label("E8", "عدد المراحل"); ws.merge_cells("F8:G8")
valuecell("F8", "=IFERROR(VLOOKUP($B$6,PartsTable,3,FALSE),\"\")")
for r in (6, 7, 8):
    ws.row_dimensions[r].height = 24

# ================= القسم 2: تفاصيل الفحص =================
ws.merge_cells("A9:G9")
s2 = ws["A9"]
s2.value = "◆  تفاصيل الفحص"
s2.font = font(13, bold=True, color=WHITE); s2.fill = fill(BLUE); s2.alignment = right
ws.row_dimensions[9].height = 26

HEAD = 10
cols = ["م", "مرحلة العمل", "عدد القطع", "الفني", "مقبول/ مرفوض", "السماحية", "السبب"]
for c, h in enumerate(cols, start=1):
    cell = ws.cell(row=HEAD, column=c, value=h)
    cell.font = font(12, bold=True, color=WHITE); cell.fill = fill(NAVY)
    cell.alignment = center; cell.border = border_box
ws.row_dimensions[HEAD].height = 30

FIRST = HEAD + 1
LAST = FIRST + MAX_STAGES - 1
for i in range(MAX_STAGES):
    r = FIRST + i
    ws.row_dimensions[r].height = 21
    c = ws.cell(row=r, column=1, value=i + 1)
    c.font = font(11, bold=True); c.alignment = center
    for col in range(2, 8):
        cc = ws.cell(row=r, column=col)
        cc.alignment = center if col in (3, 5, 6) else right
        cc.font = font(11)
    for col in range(1, 8):
        ws.cell(row=r, column=col).border = border_all

# ---------------- التنسيق الشرطي ----------------
body = f"A{FIRST}:G{LAST}"
ws.conditional_formatting.add(
    body, FormulaRule(formula=[f'AND($F$8<>"",$A{FIRST}>$F$8)'], fill=fill(GREY)))
dec = f"E{FIRST}:E{LAST}"
ws.conditional_formatting.add(dec, FormulaRule(formula=[f'$E{FIRST}="مقبول"'], fill=fill(GREEN)))
ws.conditional_formatting.add(dec, FormulaRule(formula=[f'$E{FIRST}="مرفوض"'], fill=fill(RED)))

# ---------------- القوائم المنسدلة ----------------
dv_part = DataValidation(type="list", formula1="=PartsList", allow_blank=True)
dv_part.promptTitle = "رقم القطعة"; dv_part.prompt = "اختر رقم القطعة من القائمة"
ws.add_data_validation(dv_part); dv_part.add(ws["B6"])

dv_dec = DataValidation(type="list", formula1='"مقبول,مرفوض"', allow_blank=True)
ws.add_data_validation(dv_dec); dv_dec.add(f"E{FIRST}:E{LAST}")

dv_qty = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
ws.add_data_validation(dv_qty); dv_qty.add(f"C{FIRST}:C{LAST}")

# ================= ملخص النتائج =================
SUM = LAST + 2
ws.merge_cells(start_row=SUM, start_column=1, end_row=SUM, end_column=7)
h = ws.cell(row=SUM, column=1, value="◆  ملخص نتائج الفحص")
h.font = font(13, bold=True, color=WHITE); h.fill = fill(BLUE); h.alignment = right
ws.row_dimensions[SUM].height = 26

summary = [
    ("إجمالي المراحل المفحوصة",
     f'=COUNTIF(E{FIRST}:E{LAST},"مقبول")+COUNTIF(E{FIRST}:E{LAST},"مرفوض")'),
    ("عدد المقبول", f'=COUNTIF(E{FIRST}:E{LAST},"مقبول")'),
    ("عدد المرفوض", f'=COUNTIF(E{FIRST}:E{LAST},"مرفوض")'),
    ("نسبة القبول %",
     f'=IFERROR(COUNTIF(E{FIRST}:E{LAST},"مقبول")/'
     f'(COUNTIF(E{FIRST}:E{LAST},"مقبول")+COUNTIF(E{FIRST}:E{LAST},"مرفوض")),0)'),
]
r0 = SUM + 1
for idx, (lbl, fml) in enumerate(summary):
    r = r0 + idx
    lc = ws.cell(row=r, column=1, value=lbl)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    lc.font = font(11, bold=True, color=NAVY); lc.fill = fill(LIGHT)
    lc.alignment = right; lc.border = border_all
    vc = ws.cell(row=r, column=4, value=fml)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=7)
    vc.font = font(12, bold=True); vc.alignment = center; vc.border = border_all
    if "نسبة" in lbl:
        vc.number_format = "0.0%"
    ws.row_dimensions[r].height = 22
    for col in range(1, 8):
        ws.cell(row=r, column=col).border = border_all

# ================= التواقيع =================
SIG = r0 + len(summary) + 1
# صف عناوين الجهات
titles = ["الضبط والجودة", "مدير المشروع / أو المنشأة", "الفاحص الفني"]
spans = [(1, 2), (3, 5), (6, 7)]
for (title, (c1, c2)) in zip(titles, spans):
    ws.merge_cells(start_row=SIG, start_column=c1, end_row=SIG, end_column=c2)
    cell = ws.cell(row=SIG, column=c1, value=title)
    cell.font = font(11, bold=True, color=WHITE); cell.fill = fill(NAVY)
    cell.alignment = center; cell.border = border_all
ws.row_dimensions[SIG].height = 24
# صف الاسم
for (c1, c2) in spans:
    ws.merge_cells(start_row=SIG+1, start_column=c1, end_row=SIG+1, end_column=c2)
    cell = ws.cell(row=SIG+1, column=c1, value="الاسم/ ..............................")
    cell.font = font(10, bold=True); cell.alignment = right; cell.border = border_all
ws.row_dimensions[SIG+1].height = 26
# صف التوقيع
for (c1, c2) in spans:
    ws.merge_cells(start_row=SIG+2, start_column=c1, end_row=SIG+2, end_column=c2)
    cell = ws.cell(row=SIG+2, column=c1, value="التوقيع/ ..............................")
    cell.font = font(10, bold=True); cell.alignment = right; cell.border = border_all
ws.row_dimensions[SIG+2].height = 26

# ---------------- تجميد + طباعة ----------------
ws.freeze_panes = f"A{FIRST}"
ws.print_title_rows = f"1:{HEAD}"
ws.page_setup.orientation = "portrait"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = ws.page_margins.right = 0.3
ws.page_margins.top = ws.page_margins.bottom = 0.4
ws.print_options.horizontalCentered = True
ws.print_area = f"A1:G{SIG+2}"

wb.move_sheet("استمارة الفحص", -1)
wb.active = wb.sheetnames.index("استمارة الفحص")

OUT = "/projects/sandbox/استمارة_الفحص_الفني_الشامل.xlsx"
wb.save(OUT)
print("تم الحفظ:", OUT)
print("الأوراق:", wb.sheetnames, "| صفوف الجدول:", FIRST, "->", LAST, "| التواقيع:", SIG)
