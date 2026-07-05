# -*- coding: utf-8 -*-
"""
يقرأ بيانات_الفحص.xlsx ويولّد استمارة Word (.docx) مملوءة لكل قطعة،
مطابقة لاستمارة الفحص الفني الشامل (قسم الدراسات والبحوث - ER-0001).

التشغيل:
    python3 generate_word.py
المخرجات: مجلد "مخرجات_استمارات_وورد/" يحتوي ملف Word لكل قطعة.
"""
import os
from collections import OrderedDict
from openpyxl import load_workbook

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------- إعدادات ----------------
DATA_XLSX = "/projects/sandbox/بيانات_الفحص.xlsx"
LOGO = "/projects/sandbox/logo.png"
OUT_DIR = "/projects/sandbox/مخرجات_استمارات_وورد"

NAVY = RGBColor(0x1F, 0x38, 0x64)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_HEX, LIGHT_HEX, GREEN_HEX, RED_HEX = "1F3864", "D9E1F2", "C6EFCE", "FFC7CE"

# ---------------- أدوات RTL / تنسيق ----------------
def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_table_rtl(table):
    tblPr = table._tbl.tblPr
    bidi = OxmlElement('w:bidiVisual')
    tblPr.append(bidi)

def make_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)

def style_run(run, size=11, bold=False, color=None):
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:cs'), 'Arial')
    rFonts.set(qn('w:ascii'), 'Arial')
    rFonts.set(qn('w:hAnsi'), 'Arial')
    # وسم rtl للتشغيل العربي
    rtl = OxmlElement('w:rtl')
    rtl.set(qn('w:val'), '1')
    rPr.append(rtl)

def fill_cell(cell, text, size=11, bold=False, color=None,
              align=WD_ALIGN_PARAGRAPH.RIGHT, shade=None, valign='center'):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    make_rtl(p)
    run = p.add_run("" if text is None else str(text))
    style_run(run, size=size, bold=bold, color=color)
    if shade:
        set_cell_shading(cell, shade)
    # محاذاة رأسية
    tcPr = cell._tc.get_or_add_tcPr()
    vA = OxmlElement('w:vAlign')
    vA.set(qn('w:val'), valign)
    tcPr.append(vA)
    return p

def set_col_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for idx, w in enumerate(widths_cm):
            row.cells[idx].width = Cm(w)

def add_heading_bar(doc, text):
    """شريط عنوان قسم بخلفية زرقاء داكنة."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(tbl)
    cell = tbl.rows[0].cells[0]
    fill_cell(cell, text, size=13, bold=True, color=WHITE,
              align=WD_ALIGN_PARAGRAPH.RIGHT, shade=NAVY_HEX)
    cell.width = Cm(18)
    return tbl

def no_space(paragraph):
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)

# ---------------- توليد استمارة قطعة واحدة ----------------
def build_form(part_header, stages, out_path):
    doc = Document()
    # هوامش وحجم الصفحة
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.4)
    sec.top_margin = sec.bottom_margin = Cm(1.2)

    # الخط الافتراضي
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn('w:cs'), 'Arial')

    # ============ الترويسة (شعار + جهة + رقم الاستمارة) ============
    head = doc.add_table(rows=1, cols=3)
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(head)
    set_col_widths(head, [7, 4, 7])
    # يمين: الجهة
    fill_cell(head.rows[0].cells[0], "الجمهورية اليمنية\nقسم الدراسات والبحوث",
              size=12, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.RIGHT)
    # وسط: الشعار
    mid = head.rows[0].cells[1]
    mid.text = ""
    mp = mid.paragraphs[0]
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(LOGO):
        mp.add_run().add_picture(LOGO, width=Cm(2.3))
    # يسار: رقم الاستمارة والتاريخ
    fill_cell(head.rows[0].cells[2],
              "رقم الاستمارة: ER-0001\nتاريخ الإصدار: ..../..../....\nرقم الإصدار: (1)",
              size=10, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT)

    # ============ العنوان ============
    ttl = doc.add_table(rows=1, cols=1)
    ttl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(ttl)
    ttl.rows[0].cells[0].width = Cm(18)
    fill_cell(ttl.rows[0].cells[0], "استمارة الفحص الفني الشامل",
              size=18, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER, shade=NAVY_HEX)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ============ بيانات الاستمارة الأساسية ============
    add_heading_bar(doc, "بيانات الاستمارة الأساسية")
    info = doc.add_table(rows=3, cols=4)
    info.style = 'Table Grid'
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(info)
    set_col_widths(info, [4, 5, 4, 5])
    pairs = [
        ("رقم القطعة", part_header['رقم القطعة'], "المجموعة", part_header['المجموعة']),
        ("اسم القطعة", part_header['اسم القطعة'], "التاريخ", part_header['التاريخ']),
    ]
    for r, (l1, v1, l2, v2) in enumerate(pairs):
        fill_cell(info.rows[r].cells[0], l1, bold=True, color=NAVY, shade=LIGHT_HEX)
        fill_cell(info.rows[r].cells[1], v1, bold=True)
        fill_cell(info.rows[r].cells[2], l2, bold=True, color=NAVY, shade=LIGHT_HEX)
        fill_cell(info.rows[r].cells[3], v2, bold=True)
    # اسم المشروع (صف كامل)
    fill_cell(info.rows[2].cells[0], "اسم المشروع", bold=True, color=NAVY, shade=LIGHT_HEX)
    info.rows[2].cells[1].merge(info.rows[2].cells[3])
    fill_cell(info.rows[2].cells[1], part_header['اسم المشروع'], bold=True)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ============ تفاصيل الفحص ============
    add_heading_bar(doc, "تفاصيل الفحص")
    cols = ["م", "مرحلة العمل", "عدد القطع", "الفني", "مقبول/ مرفوض", "السماحية", "السبب"]
    widths = [1.2, 4.2, 2.0, 3.0, 2.4, 2.2, 3.0]
    tbl = doc.add_table(rows=1 + len(stages), cols=len(cols))
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(tbl)
    set_col_widths(tbl, widths)
    # رأس الجدول
    for c, h in enumerate(cols):
        fill_cell(tbl.rows[0].cells[c], h, size=11, bold=True, color=WHITE,
                  align=WD_ALIGN_PARAGRAPH.CENTER, shade=NAVY_HEX)
    # الصفوف
    for i, st in enumerate(stages, start=1):
        cells = tbl.rows[i].cells
        fill_cell(cells[0], st['م'], align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(cells[1], st['مرحلة العمل'])
        fill_cell(cells[2], st['عدد القطع'], align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(cells[3], st['الفني'])
        dec = (st['مقبول/مرفوض'] or "").strip()
        shade = GREEN_HEX if dec == "مقبول" else (RED_HEX if dec == "مرفوض" else None)
        fill_cell(cells[4], dec, align=WD_ALIGN_PARAGRAPH.CENTER, shade=shade, bold=True)
        fill_cell(cells[5], st['السماحية'], align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(cells[6], st['السبب'])

    # ============ ملخص النتائج ============
    accepted = sum(1 for s in stages if (s['مقبول/مرفوض'] or "").strip() == "مقبول")
    rejected = sum(1 for s in stages if (s['مقبول/مرفوض'] or "").strip() == "مرفوض")
    total = accepted + rejected
    pct = f"{(accepted/total*100):.1f}%" if total else "0.0%"
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_heading_bar(doc, "ملخص نتائج الفحص")
    summ = doc.add_table(rows=1, cols=4)
    summ.style = 'Table Grid'
    summ.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(summ)
    set_col_widths(summ, [4.5, 4.5, 4.5, 4.5])
    labels = [f"الإجمالي: {total}", f"مقبول: {accepted}",
              f"مرفوض: {rejected}", f"نسبة القبول: {pct}"]
    for c, lab in enumerate(labels):
        fill_cell(summ.rows[0].cells[c], lab, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER, shade=LIGHT_HEX)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ============ التواقيع ============
    sig = doc.add_table(rows=3, cols=3)
    sig.style = 'Table Grid'
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(sig)
    set_col_widths(sig, [6, 6, 6])
    titles = ["الضبط والجودة", "مدير المشروع / أو المنشأة", "الفاحص الفني"]
    for c, t in enumerate(titles):
        fill_cell(sig.rows[0].cells[c], t, bold=True, color=WHITE,
                  align=WD_ALIGN_PARAGRAPH.CENTER, shade=NAVY_HEX)
    for c in range(3):
        fill_cell(sig.rows[1].cells[c], "الاسم/ ..........................", size=10)
        fill_cell(sig.rows[2].cells[c], "التوقيع/ ........................", size=10)

    doc.save(out_path)

# ---------------- القراءة والتجميع ----------------
def read_data():
    wb = load_workbook(DATA_XLSX, data_only=True)
    ws = wb["بيانات الفحص"]
    headers = [c.value for c in ws[2]]  # الصف 2 = رأس الأعمدة
    idx = {h: i for i, h in enumerate(headers)}
    groups = OrderedDict()
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        part = row[idx["رقم القطعة"]]
        if not part:
            continue
        rec = {h: (row[idx[h]] if row[idx[h]] is not None else "") for h in headers}
        groups.setdefault(str(part), []).append(rec)
    return groups

def safe_name(s):
    return "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in str(s)).strip()

def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    groups = read_data()
    if not groups:
        print("لا توجد بيانات في الملف.")
        return
    made = []
    for part, rows in groups.items():
        first = rows[0]
        header = {
            'رقم القطعة': first['رقم القطعة'],
            'اسم القطعة': first['اسم القطعة'],
            'اسم المشروع': first['اسم المشروع'],
            'التاريخ': first['التاريخ'],
            'المجموعة': first['المجموعة'],
        }
        stages = sorted(rows, key=lambda r: (r['م'] if isinstance(r['م'], (int, float)) else 0))
        out = os.path.join(OUT_DIR, f"استمارة_فحص_{safe_name(part)}.docx")
        build_form(header, stages, out)
        made.append((part, len(stages), out))
    print(f"تم توليد {len(made)} استمارة Word في: {OUT_DIR}")
    for part, n, path in made:
        print(f"  - {part}: {n} مرحلة -> {os.path.basename(path)}")

if __name__ == "__main__":
    main()
